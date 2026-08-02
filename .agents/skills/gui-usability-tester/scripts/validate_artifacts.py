"""Validate portable GUI/usability skill artifacts that use the output contract."""

from __future__ import annotations

import csv
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


ALLOWED_STATUS = {"Pass", "Fail", "Passed", "Failed", "Blocked", "Not Executed"}
ALIASES = {
    "check_id": ("check_id", "Item ID"),
    "requirement_id": ("requirement_id", "FR ID"),
    "interface_aspect": ("interface_aspect", "Interface Aspect (IA)"),
    "screen": ("screen", "Screen"),
    "precondition": ("precondition", "Preconditions"),
    "action": ("action", "Test Item Description"),
    "expected_result": ("expected_result", "Expected Result"),
    "actual_result": ("actual_result", "Actual Result"),
    "status": ("status", "Status"),
    "environment": ("environment", "Environment"),
    "evidence_ref": ("evidence_ref", "Evidence Reference"),
    "defect_id": ("defect_id", "Bug ID"),
    "origin": ("origin", "Origin"),
    "notes": ("notes", "Notes"),
}


def read_rows(path: Path) -> tuple[set[str], list[dict[str, str]]]:
    if path.suffix.lower() == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=False)
        sheet = workbook["GUI Checklist"] if "GUI Checklist" in workbook.sheetnames else workbook.active
        values = sheet.iter_rows(values_only=True)
        headers = [str(item).strip() if item is not None else "" for item in next(values, ())]
        rows = [dict(zip(headers, ("" if item is None else item for item in row))) for row in values]
        return set(headers), rows
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        return set(reader.fieldnames or []), list(reader)


def value(row: dict[str, str], canonical: str) -> str:
    return next(((row.get(name) or "").strip() for name in ALIASES[canonical] if name in row), "")


def main(root_arg: str) -> int:
    root = Path(root_arg).resolve()
    checklist_dir = root / "checklist"
    checklist = checklist_dir / "gui_checklist.xlsx"
    if not checklist.is_file():
        checklist = checklist_dir / "gui_checklist.csv"
    defects = root / "bugs" / "bug_report.md"
    errors: list[str] = []
    warnings: list[str] = []
    status_counts: Counter[str] = Counter()
    failed_links: list[tuple[str, str, str]] = []

    if not checklist.is_file():
        errors.append(f"Missing checklist: {checklist}")
    else:
        fields, rows = read_rows(checklist)
        missing = {key for key, names in ALIASES.items() if not any(name in fields for name in names)}
        if missing:
            errors.append("Checklist missing columns: " + ", ".join(sorted(missing)))
        ids: set[str] = set()
        for number, row in enumerate(rows, start=2):
            if None in row:
                errors.append(f"Checklist row {number}: contains {len(row[None] or [])} extra value(s)")
            missing_values = [name for name, raw in row.items() if name is not None and raw is None]
            if missing_values:
                errors.append(
                    f"Checklist row {number}: has fewer values than the header; missing "
                    + ", ".join(missing_values)
                )
            check_id = value(row, "check_id")
            status = value(row, "status")
            status_counts[status] += 1
            if not check_id:
                errors.append(f"Checklist row {number}: missing check_id")
            elif check_id in ids:
                errors.append(f"Checklist row {number}: duplicate check_id {check_id}")
            ids.add(check_id)
            if status not in ALLOWED_STATUS:
                errors.append(f"Checklist {check_id or number}: invalid status {status!r}")
            if status in {"Fail", "Failed"} and value(row, "defect_id") in {"", "N/A"}:
                errors.append(f"Checklist {check_id or number}: Fail without defect_id")
            if status in {"Fail", "Failed"}:
                failed_links.append((check_id, value(row, "defect_id"), value(row, "evidence_ref")))
            if status in {"Pass", "Fail", "Passed", "Failed"} and not value(row, "actual_result"):
                errors.append(f"Checklist {check_id or number}: executed without actual_result")
            if status in {"Blocked", "Not Executed"} and not value(row, "actual_result"):
                warnings.append(f"Checklist {check_id or number}: {status} has no reason in actual_result")

    if not defects.is_file():
        errors.append(f"Missing defect report: {defects}")
    else:
        defect_text = defects.read_text(encoding="utf-8-sig")
        for check_id, defect_id, evidence_ref in failed_links:
            if defect_id and not re.search(rf"(?<![A-Z0-9-]){re.escape(defect_id)}(?![A-Z0-9-])", defect_text):
                errors.append(f"Checklist {check_id}: defect {defect_id} absent from defect report")
            if not evidence_ref or evidence_ref == "N/A":
                errors.append(f"Checklist {check_id}: failure has no evidence reference")
            else:
                for raw_ref in re.split(r"[;,]", evidence_ref):
                    ref = raw_ref.strip().strip("`")
                    if ref and "://" not in ref and not (root / Path(ref)).is_file():
                        errors.append(f"Checklist {check_id}: missing evidence file {ref}")

    participant_list = root / "usability" / "participant_list.md"
    session_results = root / "usability" / "usability_results.xlsx"
    if participant_list.exists() != session_results.exists():
        warnings.append("Usability participant register and session results are not both present")

    if errors:
        print("VALIDATION FAILED")
        for error in errors:
            print(f"- {error}")
        return 1

    print("VALIDATION PASSED")
    print("Checklist statuses: " + ", ".join(f"{key}={count}" for key, count in sorted(status_counts.items())))
    print(f"Failed checks reconciled: {len(failed_links)}")
    for warning in warnings:
        print(f"WARNING: {warning}")
    return 0


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("Usage: validate_artifacts.py <deliverables-root>")
    raise SystemExit(main(sys.argv[1]))
