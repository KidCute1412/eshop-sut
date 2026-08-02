#!/usr/bin/env python3
"""Export a gui-usability checklist.md table to a real .xlsx file for the required Excel deliverable."""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from _table_utils import ensure_utf8_stdio, parse_table

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    Workbook = None

HEADER = [
    "No.",
    "Type",
    "Checkpoint",
    "IA",
    "Requirement / Heuristic Reference",
    "Platform",
    "Source",
    "AI-Miss Reason",
    "Yes",
    "No",
    "Remarks",
    "Evidence",
    "Bug ID",
]

COLUMN_WIDTHS = [8, 10, 55, 8, 30, 12, 14, 30, 6, 6, 55, 30, 14]


def main() -> int:
    ensure_utf8_stdio()
    if Workbook is None:
        print("ERROR: openpyxl is not installed. Run: pip install openpyxl", file=sys.stderr)
        return 3

    parser = argparse.ArgumentParser(description="Convert checklist.md into a real .xlsx file (open in Excel).")
    parser.add_argument("--input", required=True, help="Path to checklist.md.")
    parser.add_argument("--output", required=True, help="Path to write the .xlsx file.")
    parser.add_argument("--sheet-name", default="GUI Checklist", help="Worksheet name.")
    args = parser.parse_args()

    input_path = Path(args.input)
    if not input_path.is_file():
        print(f"ERROR: input file not found: {input_path}", file=sys.stderr)
        return 2

    rows = parse_table(input_path.read_text(encoding="utf-8"), HEADER)
    if not rows:
        print(f"ERROR: zero checklist rows parsed from {input_path}", file=sys.stderr)
        return 2

    wb = Workbook()
    ws = wb.active
    ws.title = args.sheet_name[:31]

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    section_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    wrap = Alignment(wrap_text=True, vertical="top")

    for col_idx, name in enumerate(HEADER, start=1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

    item_count = 0
    for r_idx, row in enumerate(rows, start=2):
        is_section = row.get("Type", "").strip() == "Section"
        if not is_section:
            item_count += 1
        for c_idx, name in enumerate(HEADER, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=row.get(name, ""))
            cell.alignment = wrap
            if is_section:
                cell.fill = section_fill
                cell.font = Font(bold=True)

    for c_idx, width in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(c_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    summary_ws = wb.create_sheet("Test Summary")
    passed = sum(1 for r in rows if r.get("Type") == "Item" and r.get("Yes", "").strip().upper() == "X")
    failed = sum(1 for r in rows if r.get("Type") == "Item" and r.get("No", "").strip().upper() == "X")
    summary_rows = [
        ("Total checklist items", item_count),
        ("Passed", passed),
        ("Failed", failed),
        ("Not executed", item_count - passed - failed),
    ]
    summary_ws.cell(row=1, column=1, value="Metric").font = Font(bold=True)
    summary_ws.cell(row=1, column=2, value="Value").font = Font(bold=True)
    for i, (label, value) in enumerate(summary_rows, start=2):
        summary_ws.cell(row=i, column=1, value=label)
        summary_ws.cell(row=i, column=2, value=value)
    summary_ws.column_dimensions["A"].width = 25
    summary_ws.column_dimensions["B"].width = 12

    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    print(f"Wrote {item_count} item row(s) ({passed} passed, {failed} failed) to {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
