"""Validate internal consistency of the current HW03 deliverables."""

from __future__ import annotations

import csv
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[4]
DELIVERABLES = ROOT / "docs/assignments/HW03/deliverables"
DATA = ROOT / "docs/assignments/HW03/workbench/data"
errors: list[str] = []


def require(condition: bool, message: str) -> None:
    if not condition:
        errors.append(message)


with (DATA / "gui_checklist.csv").open(
    encoding="utf-8-sig", newline=""
) as source:
    rows = list(csv.DictReader(source))

require(len(rows) == 45, f"Checklist must contain 45 rows; found {len(rows)}")
counts = Counter(row["Status"] for row in rows)
require(counts == {"Passed": 25, "Failed": 20}, f"Unexpected status counts: {counts}")
require(all(row["Actual Result"].strip() for row in rows), "Every checklist row needs an actual result")
require(all(row["Notes"].strip() for row in rows), "Every checklist row needs execution or non-execution notes")
require(all(row["Screen"].strip() for row in rows), "Every checklist row needs a screen")
require(all(row["Environment"].strip() for row in rows), "Every checklist row needs an environment")
require(all(row["Origin"] in {"AI", "Human-added", "Hybrid"} for row in rows), "Every checklist row needs a valid origin")
require(any(row["Origin"] == "AI" for row in rows), "Checklist must identify AI-originated items")
require(any(row["Origin"] == "Human-added" for row in rows), "Checklist must identify human-added items")
require(all(row["AI Critique / Reason Missed"].strip() for row in rows), "Every checklist row needs an AI-review rationale")
require(all("initial AI" in row["AI Critique / Reason Missed"] or row["Origin"] != "Human-added" for row in rows), "Every human-added item must explain the initial AI omission")
allowed_fr_ids = {"FR-07", "FR-10", "FR-11", "FR-18", "FR-23"}
require(set(row["FR ID"] for row in rows) <= allowed_fr_ids, "Checklist contains an out-of-scope FR ID")
require(sum(row["FR ID"] == "FR-23" for row in rows) == 14, "Expected exactly fourteen FR-23 rows")
require(all(row["FR ID"] == "FR-23" for row in rows if "Mobile" in row["Screen"]), "Every Mobile screen must map to FR-23")
require(not any("Mobile" in row["Screen"] and row["FR ID"] in {"FR-07", "FR-11"} for row in rows), "Mobile scope remains on FR-07/FR-11")
require(Counter(row["Status"] for row in rows if row["FR ID"] == "FR-23") == {"Passed": 9, "Failed": 5}, "Unexpected FR-23 results")
require(all("Expo Go" in row["Environment"] for row in rows if row["FR ID"] == "FR-23"), "Every classified FR-23 result must identify the Mobile execution environment")
require(all(row["Evidence Reference"].strip() for row in rows if row["FR ID"] == "FR-23"), "Every FR-23 result must reference runtime evidence")
require(all("Static source review" not in row["Environment"] and "Source-derived" not in row["Notes"] for row in rows), "No Passed/Failed checklist result may rely on static source review")

failed_ids: set[str] = set()
for row in rows:
    item_id = row["Item ID"]
    if row["Status"] == "Failed":
        require(row["Bug ID"].startswith("BUG-"), f"{item_id} lacks Bug ID")
        require(bool(row["Evidence Reference"]), f"{item_id} lacks evidence reference")
        if row["Evidence Reference"]:
            evidence_refs = [ref.strip() for ref in row["Evidence Reference"].split(";") if ref.strip()]
            require(all((DELIVERABLES / ref).is_file() for ref in evidence_refs), f"Missing evidence for {item_id}")
        failed_ids.add(row["Bug ID"])
    elif row["Status"] == "Passed":
        require(row["Bug ID"] == "N/A", f"Passed {item_id} must not map to a bug")
        if row["Evidence Reference"]:
            evidence_refs = [ref.strip() for ref in row["Evidence Reference"].split(";") if ref.strip()]
            require(all((DELIVERABLES / ref).is_file() for ref in evidence_refs), f"Missing passed-item evidence for {item_id}")
    else:
        require("Mobile screenshot" in row["Notes"] or "physical device" in row["Notes"] or "Mobile app" in row["Notes"], f"{item_id} lacks a qualifying Mobile reason")

bug_report = (DELIVERABLES / "bugs/bug_report.md").read_text(encoding="utf-8")
verified_headings = set(re.findall(r"^### (BUG-\d{3}) —", bug_report, flags=re.MULTILINE))
require(verified_headings == failed_ids, f"Bug headings differ from failed rows: {verified_headings ^ failed_ids}")
require("Verified defects | 19 runtime-observed defects" in bug_report, "Bug summary count missing")
require("BUG-024" not in bug_report, "Retired BUG-024 remains in bug report")

bug_pngs = sorted((DELIVERABLES / "bugs/evidence_images").glob("*.png"))
chrome_pngs = sorted((DELIVERABLES / "cross_platform/chrome_desktop").glob("*.png"))
firefox_pngs = sorted((DELIVERABLES / "cross_platform/firefox_desktop").glob("*.png"))
mobile_pngs = sorted((DELIVERABLES / "cross_platform/mobile_real_device").glob("*.png"))
require(len(bug_pngs) == 15, f"Expected 15 bug PNGs; found {len(bug_pngs)}")
require(len(chrome_pngs) == 5, f"Expected 5 Chrome PNGs; found {len(chrome_pngs)}")
require(len(firefox_pngs) == 5, f"Expected 5 Firefox PNGs; found {len(firefox_pngs)}")
require(len(mobile_pngs) == 4, f"Expected four supplied Mobile screenshots; found {len(mobile_pngs)}")

script = (DELIVERABLES / "usability/usability_test_script.md").read_text(encoding="utf-8")
findings = (DELIVERABLES / "usability/usability_findings.md").read_text(encoding="utf-8")
recording_index = (DELIVERABLES / "usability/recordings/video_links.md").read_text(encoding="utf-8")
for expected in ("Opening and task statement", "Step-by-step execution", "intervention", "Post-task questionnaire", "SUS"):
    require(expected in script, f"Usability test script missing: {expected}")
require("not participant observations" in findings, "Findings need an evidence-integrity warning")
shared_recording_url = "https://drive.google.com/drive/u/0/folders/1_3wIHUVqJGG-mPwcZotoAStgRjd6X9x_"
require(shared_recording_url in recording_index, "Usability recording index lacks the supplied shared Drive folder")
require(all(session_id in recording_index for session_id in ("Pilot", "P1", "P2", "P3", "P4", "P5", "P6", "P7")), "Recording index lacks a required session")

with (DATA / "session_results.csv").open(encoding="utf-8-sig", newline="") as source:
    session_rows = list(csv.DictReader(source))
require([row["session_id"] for row in session_rows] == ["Pilot", "P1", "P2", "P3", "P4", "P5", "P6", "P7"], "Unexpected usability session IDs")
for row in session_rows:
    if row["evidence_status"] == "Complete":
        require(bool(row["recording_ref"].strip()), f"{row['session_id']} is complete without recording evidence")
        require(bool(row["completion"].strip()), f"{row['session_id']} is complete without a completion result")
        require(bool(row["duration_seconds"].strip()), f"{row['session_id']} is complete without task duration")

with (DATA / "sus_survey_results.csv").open(encoding="utf-8-sig", newline="") as source:
    sus_rows = list(csv.DictReader(source))
require([row["participant_id"] for row in sus_rows] == ["P1", "P2", "P3", "P4", "P5", "P6", "P7", "AVERAGE"], "Unexpected SUS participant IDs")
for row in sus_rows[:-1]:
    responses = [row[f"q{number}"].strip() for number in range(1, 11)]
    require(not any(responses) or all(responses), f"{row['participant_id']} has a partial SUS response")
    if all(responses):
        require(all(value in {"1", "2", "3", "4", "5"} for value in responses), f"{row['participant_id']} has an invalid SUS value")
        require(bool(row["sus_score"].strip()), f"{row['participant_id']} has responses but no SUS score")

readme = (DELIVERABLES / "README.md").read_text(encoding="utf-8")
for expected in ("| Checklist items executed | 45 (31 desktop/API; 14 FR-23, including screenshot-backed Mobile checks) |", "| Passed | 25 |", "| Failed | 20 |", "| Not executed | 0 checklist items; CP-03 metadata/overlay completion remains pending |", "| Verified defects | 19 runtime-observed; four Mobile screenshots still need the required identity overlay |"):
    require(expected in readme, f"README missing summary: {expected}")

critique = (DELIVERABLES / "ai_reports/ai_critique.md").read_text(encoding="utf-8")
critique_words = len(re.findall(r"\b[\w’'-]+\b", re.sub(r"^#.*$", "", critique, flags=re.MULTILINE), flags=re.UNICODE))
require(200 <= critique_words <= 300, f"AI Critique has {critique_words} words")

for path in DELIVERABLES.rglob("*"):
    if path.is_file() and path.suffix.lower() in {".md", ".csv", ".txt"}:
        try:
            path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError as exc:
            errors.append(f"UTF-8 error in {path.relative_to(DELIVERABLES)}: {exc}")

for rel, rows_expected, cols_expected in (
    ("checklist/gui_checklist.xlsx", 46, 15),
):
    workbook = load_workbook(DELIVERABLES / rel, read_only=False)
    require(workbook.sheetnames == ["Test Summary", "GUI Checklist"], f"{rel} must contain Test Summary and GUI Checklist")
    sheet = workbook["GUI Checklist"]
    require(sheet.max_row == rows_expected, f"{rel} row count is {sheet.max_row}")
    require(sheet.max_column == cols_expected, f"{rel} column count is {sheet.max_column}")
    require(sheet.freeze_panes == "A2", f"{rel} should freeze A2")
    require(bool(sheet.auto_filter.ref), f"{rel} lacks an auto-filter")
    require(bool(sheet.print_area), f"{rel} lacks a print area")
    summary = workbook["Test Summary"]
    summary_values = {summary.cell(row=i, column=1).value: summary.cell(row=i, column=2).value for i in range(1, summary.max_row + 1)}
    require("COUNTA" in str(summary_values.get("Designed")), "XLSX summary Designed must be formula-derived")
    require("COUNTIF" in str(summary_values.get("Executed")), "XLSX summary Executed must be formula-derived")
    require("COUNTIF" in str(summary_values.get("Passed")), "XLSX summary Passed must be formula-derived")
    require("COUNTIF" in str(summary_values.get("Failed")), "XLSX summary Failed must be formula-derived")
    require("COUNTIF" in str(summary_values.get("Not Executed")), "XLSX summary Not Executed must be formula-derived")
    require(len(sheet.data_validations.dataValidation) == 3, "XLSX checklist must validate FR, Status, and Origin values")

for rel in ("README.md", "main_report.md"):
    report_text = (DELIVERABLES / rel).read_text(encoding="utf-8")
    require("23127404@student.hcmus.edu.vn" in report_text, f"{rel} lacks the real contact email")
    require("23127404@hcmus.edu.vn" in report_text, f"{rel} lacks the required overlay identity")
    require("Pool D" not in report_text, f"{rel} still contains Pool D")

usability_workbook = load_workbook(DELIVERABLES / "usability/usability_results.xlsx", data_only=False)
require(usability_workbook.sheetnames == ["Sessions", "Observations", "SUS", "Recordings"], "Unexpected usability workbook sheets")
require(usability_workbook["Sessions"].max_row == 9, "Usability workbook needs Pilot and P1-P7")
require(usability_workbook["Observations"].max_row == 65, "Usability workbook needs eight checkpoints for eight sessions")
require(usability_workbook["Recordings"].max_row == 9, "Recording index needs Pilot and P1-P7")
recordings_sheet = usability_workbook["Recordings"]
expected_recording_headers = ["session_id", "filename", "drive_folder_url", "duration_seconds", "file_size_bytes", "sha256", "resolution", "local_verified", "drive_access_verified", "status"]
require([cell.value for cell in recordings_sheet[1]] == expected_recording_headers, "Unexpected Recordings sheet schema")
require(all(recordings_sheet.cell(row=row, column=3).value == shared_recording_url for row in range(2, 10)), "Usability workbook does not reference the shared Drive folder for every session")
require([recordings_sheet.cell(row=row, column=2).value for row in range(2, 10)] == ["Pilot.mp4", "P1.mp4", "P2.mp4", "P3.mp4", "P4.mp4", "P5.mp4", "P6.mp4", "P7.mp4"], "Recording filenames do not map Pilot and P1-P7")
require(all(recordings_sheet.cell(row=row, column=8).value == "Yes" for row in range(2, 10)), "Local recording verification is incomplete")
require(all(len(str(recordings_sheet.cell(row=row, column=6).value or "")) == 64 for row in range(2, 10)), "Recording SHA-256 evidence is incomplete")
for usability_sheet in usability_workbook.worksheets:
    require(usability_sheet.freeze_panes == "A2", f"{usability_sheet.title} should freeze A2")
    require(bool(usability_sheet.auto_filter.ref), f"{usability_sheet.title} lacks an auto-filter")
sus_sheet = usability_workbook["SUS"]
require(sus_sheet.max_row == 9 and sus_sheet.max_column == 18, "Unexpected SUS sheet dimensions")
require(len(sus_sheet.data_validations.dataValidation) == 1, "SUS workbook lacks 1–5 response validation")
require(str(sus_sheet["L2"].value).startswith("=IF(COUNTA"), "SUS workbook lacks adjusted-score formula")
require(str(sus_sheet["M2"].value).startswith("=IF(ISNUMBER"), "SUS workbook lacks participant-score formula")
require("AVERAGE" in str(sus_sheet["M9"].value), "SUS workbook lacks aggregate formula")

for rel in ("main_report.pdf", "ai_reports/ai_audit_report.pdf", "ai_reports/ai_critique.pdf"):
    reader = PdfReader(DELIVERABLES / rel)
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    require("23127404" in text, f"{rel} lacks student ID")
    require("|---" not in text, f"{rel} contains raw Markdown table syntax")
    require(reader.metadata and reader.metadata.title, f"{rel} lacks title metadata")

if errors:
    print("VALIDATION FAILED")
    for error in errors:
        print(f"- {error}")
    raise SystemExit(1)

print("VALIDATION PASSED")
print(f"Checklist: {dict(counts)}")
print(f"Verified bugs/screenshots: {len(failed_ids)}/{len(bug_pngs)}")
print(f"Chrome/Firefox/Mobile screenshots: {len(chrome_pngs)}/{len(firefox_pngs)}/{len(mobile_pngs)}")
print(f"AI Critique words: {critique_words}")
